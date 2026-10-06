"""Convert research/paper/manuscript.md (a restricted Markdown subset) to a formatted .docx.
Supported: '# ' title, '## ' / '### ' headings, paragraphs, '> ' display equations, '- ' bullets,
'1. ' numbered items, pipe tables, '![..](path)' images, '*Figure ...*' captions, **bold**, *italic*.
"""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC, OUT = sys.argv[1], sys.argv[2]
BASE = SRC.rsplit("/", 1)[0] + "/"
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
for s in ("Heading 1", "Heading 2", "Title"):
    doc.styles[s].font.name = "Times New Roman"; doc.styles[s].font.color.rgb = RGBColor(0, 0, 0)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(1)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


lines = open(SRC, encoding="utf8").read().split("\n")
i = 0; para_buf = []


def flush():
    global para_buf
    if para_buf:
        p = doc.add_paragraph(); add_inline(p, " ".join(s.strip() for s in para_buf))
        p.paragraph_format.space_after = Pt(6); para_buf = []


while i < len(lines):
    ln = lines[i]
    if ln.startswith("# "):
        flush(); t = doc.add_paragraph(ln[2:].strip(), style="Title"); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("## "):
        flush(); doc.add_heading(ln[3:].strip(), level=1)
    elif ln.startswith("### "):
        flush(); doc.add_heading(ln[4:].strip(), level=2)
    elif ln.strip() == "---" or ln.strip() == "":
        flush()
    elif ln.startswith("> "):
        flush(); p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5)
        r = p.add_run(ln[2:].strip()); r.font.name = "Cambria Math"; r.italic = False
    elif ln.startswith("- "):
        flush(); p = doc.add_paragraph(style="List Bullet"); add_inline(p, ln[2:].strip())
    elif re.match(r"^\d+\. ", ln):
        flush(); p = doc.add_paragraph(style="List Number"); add_inline(p, re.sub(r"^\d+\. ", "", ln).strip())
    elif re.match(r"^\(?[ivx]+\)\s", ln) or re.match(r"^\[\d+\] ", ln):
        flush(); p = doc.add_paragraph(); add_inline(p, ln.strip()); p.paragraph_format.space_after = Pt(3)
    elif ln.startswith("!["):
        flush(); path = re.search(r"\((.+?)\)", ln).group(1)
        doc.add_picture(BASE + path, width=Inches(6.3)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("*Figure") or ln.startswith("*Table"):
        flush(); p = doc.add_paragraph(); r = p.add_run(ln.strip("*")); r.italic = True; r.font.size = Pt(9.5)
    elif ln.startswith("|"):
        flush(); rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                rows.append(cells)
            i += 1
        ncol = max(len(r) for r in rows)
        tb = doc.add_table(rows=len(rows), cols=ncol); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, r in enumerate(rows):
            for ci in range(ncol):
                cell = tb.cell(ri, ci); cell.text = ""
                p = cell.paragraphs[0]; add_inline(p, r[ci] if ci < len(r) else "")
                for run in p.runs:
                    run.font.size = Pt(9); run.bold = run.bold or ri == 0
        doc.add_paragraph()
        continue
    else:
        para_buf.append(ln)
    i += 1
flush()
doc.save(OUT)
print("saved", OUT)
